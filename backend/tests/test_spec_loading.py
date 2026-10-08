"""The Excel file is the source of truth; plans are validated against it and the tool registry."""
import shutil

import openpyxl
import yaml

from flowline.spec import loader
from flowline.spec.loader import read_excel_specs

from .conftest import ROOT


def test_excel_defines_ten_workflows(settings):
    specs, problems = read_excel_specs(settings.workflow_file)
    assert problems == []
    assert [s.id for s in specs] == [f"WF{n:03d}" for n in range(1, 11)]
    wf1 = specs[0]
    assert wf1.steps == ["Load inventory", "compare current stock with minimum threshold", "identify low-stock products",
                         "calculate reorder quantity", "generate restock list"]
    assert wf1.tools == ["CSV reader", "calculator"]
    assert all(len(s.test_requests) == 1 for s in specs)


def test_every_plan_is_valid_and_covers_every_excel_step(settings):
    reg = loader.get_registry()
    assert reg.problems == []
    for wf in reg.all():
        assert wf.executable, (wf.id, wf.issues)
        assert wf.warnings == [], (wf.id, wf.warnings)   # no Excel step left unmapped


def test_decision_thresholds_are_read_from_excel(settings):
    reg = loader.get_registry()
    assert reg.get("WF002").params["threshold_pct"] == 10
    assert reg.get("WF002").param_sources["threshold_pct"] == "Excel decision logic"
    assert reg.get("WF010").params["max_failure_rate"] == 10


def test_changing_the_excel_threshold_changes_behaviour(tmp_path, settings):
    xlsx = tmp_path / "wf.xlsx"
    shutil.copy(settings.workflow_file, xlsx)
    wb = openpyxl.load_workbook(xlsx)
    ws = wb["Workflows"]
    for row in ws.iter_rows(min_row=2):
        if row[0].value == "WF002":
            row[5].value = "Flag when price difference exceeds 20%"
    wb.save(xlsx)
    settings.workflow_file = xlsx
    reg = loader.reset_registry(settings)
    assert reg.get("WF002").params["threshold_pct"] == 20


def test_headers_are_matched_tolerantly(tmp_path, settings):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "workflows"
    ws.append(["workflow id", "Workflow Name", "STEPS", "trigger"])
    ws.append(["X1", "Demo", "a -> b -> c", "when asked"])
    p = tmp_path / "x.xlsx"
    wb.save(p)
    specs, problems = read_excel_specs(p)
    assert problems == [] and specs[0].steps == ["a", "b", "c"] and specs[0].trigger == "when asked"


def test_broken_plan_is_reported_not_crashing(tmp_path, settings):
    plans = tmp_path / "plans"
    shutil.copytree(settings.plans_dir, plans)
    raw = yaml.safe_load((plans / "WF001.yaml").read_text(encoding="utf-8"))
    raw["steps"][0]["tool"] = "load_tabel"                    # typo in tool name
    step = next(st for st in raw["steps"] if st["id"] == "threshold")
    step["args"]["colums"] = step["args"].pop("columns")       # typo in arg name
    (plans / "WF001.yaml").write_text(yaml.safe_dump(raw), encoding="utf-8")
    settings.plans_dir = plans
    reg = loader.reset_registry(settings)
    wf = reg.get("WF001")
    assert not wf.executable
    assert any("unknown tool 'load_tabel'" in i for i in wf.issues)
    assert any("no argument(s) colums" in i for i in wf.issues)
    assert reg.get("WF002").executable                        # other workflows unaffected


def test_hot_reload_picks_up_new_plan(tmp_path, settings):
    plans = tmp_path / "plans"
    shutil.copytree(settings.plans_dir, plans)
    (plans / "WF001.yaml").unlink()
    settings.plans_dir = plans
    reg = loader.reset_registry(settings)
    assert not reg.get("WF001").executable
    shutil.copy(ROOT / "workflows" / "plans" / "WF001.yaml", plans / "WF001.yaml")
    assert reg.get("WF001").executable                        # ensure_fresh() reloads on change