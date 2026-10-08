"""Demo helper: add the WF011 row to the Excel file and install its plan.

    python scripts/add_wf011.py            # add row + copy plan
    python scripts/add_wf011.py --remove   # undo (restore the original 10 workflows)

In the Loom you can instead add the row by hand in Excel — this script does exactly that.
"""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "workflows" / "AI_Agent_Workflow_Assessment_1.xlsx"
PLAN_SRC = ROOT / "examples" / "wf011" / "WF011.yaml"
PLAN_DST = ROOT / "workflows" / "plans" / "WF011.yaml"

ROW = {
    "Workflow_ID": "WF011",
    "Workflow_Name": "Low Margin Alert",
    "Trigger": "User asks which products have low profit margins",
    "Inputs": "Product CSV with prices and unit costs; minimum margin",
    "Steps": "Load products → calculate gross margin → flag low-margin products → rank by margin → generate alert list",
    "Decision_Logic": "Flag products whose gross margin is below 62%",
    "Tools_Required": "CSV reader; calculator",
    "Expected_Output": "List of low-margin products with price, cost, margin and gap to target",
}
TEST = ("WF011", "Which products have low margins?", "Tests a workflow added without code changes")


def add() -> None:
    backup = XLSX.with_suffix(".xlsx.bak")
    if not backup.exists():
        shutil.copy(XLSX, backup)  # original kept so --remove can restore it byte-for-byte
    wb = openpyxl.load_workbook(XLSX)
    ws = wb["Workflows"]
    headers = [c.value for c in ws[1]]
    if any(r[0] == "WF011" for r in ws.iter_rows(min_row=2, values_only=True)):
        print("WF011 already present in Excel")
    else:
        ws.append([ROW.get(h, "") for h in headers])
        if "Test_Questions" in wb.sheetnames:
            wb["Test_Questions"].append(list(TEST))
        wb.save(XLSX)
        print(f"Added WF011 to {XLSX.name}")
    shutil.copy(PLAN_SRC, PLAN_DST)
    print(f"Installed plan {PLAN_DST.relative_to(ROOT)}")


def remove() -> None:
    backup = XLSX.with_suffix(".xlsx.bak")
    if backup.exists():
        shutil.move(backup, XLSX)
        PLAN_DST.unlink(missing_ok=True)
        print("Restored the original Excel file and removed the WF011 plan")
        return
    wb = openpyxl.load_workbook(XLSX)
    for name in ("Workflows", "Test_Questions"):
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        for idx in range(ws.max_row, 1, -1):
            if ws.cell(idx, 1).value == "WF011":
                ws.delete_rows(idx)
    wb.save(XLSX)
    PLAN_DST.unlink(missing_ok=True)
    print("Removed WF011 from Excel and plans")


if __name__ == "__main__":
    remove() if "--remove" in sys.argv else add()